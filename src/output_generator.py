"""Generate output files: JSON and Excel."""

from __future__ import annotations

import json
import os
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def save_json(data: dict[str, Any], path: str) -> None:
    """Save results to a JSON file."""
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def save_excel(
    calendar_data: list[dict[str, Any]],
    brief_data: list[dict[str, Any]],
    path: str,
) -> None:
    """Save results to an Excel file with two sheets."""
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb = Workbook()

    # --- Sheet 1: PR Calendar ---
    ws_cal = wb.active
    ws_cal.title = "PR-календарь"
    _write_calendar_sheet(ws_cal, calendar_data)

    # --- Sheet 2: Strategic PR Brief ---
    ws_brief = wb.create_sheet("Стратегический PR-бриф")
    _write_brief_sheet(ws_brief, brief_data)

    wb.save(path)


# Header style
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_WRAP = Alignment(wrap_text=True, vertical="top")

# Status colors
_STATUS_FILLS = {
    "accepted": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "review": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "rejected": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
}

CALENDAR_COLUMNS = [
    ("project_id", "Проект", 15),
    ("date", "Дата события", 14),
    ("title", "Название события", 40),
    ("event_type", "Тип события", 16),
    ("pr_score", "PR Score", 10),
    ("why_it_matters", "Почему это инфоповод", 45),
    ("how_to_strengthen", "Как усилить", 35),
    ("pr_format", "Формат PR-отработки", 25),
    ("quality_status", "Статус обработки", 16),
    ("quality_comment", "Комментарий QC", 35),
]

BRIEF_COLUMNS = [
    ("project_id", "Проект", 15),
    ("title", "Краткое название", 35),
    ("category", "Категория возможности", 22),
    ("media_potential", "Описание медиапотенциала", 45),
    ("timing", "Тайминг", 20),
    ("effort_level", "Уровень усилий", 14),
    ("media_formats", "Рекомендуемые форматы", 30),
    ("recommendation", "Рекомендация", 45),
    ("quality_status", "Статус обработки", 16),
    ("quality_comment", "Комментарий QC", 35),
]


def _write_calendar_sheet(ws, data: list[dict]) -> None:
    _write_sheet(ws, data, CALENDAR_COLUMNS)


def _write_brief_sheet(ws, data: list[dict]) -> None:
    _write_sheet(ws, data, BRIEF_COLUMNS)


def _write_sheet(ws, data: list[dict], columns: list[tuple[str, str, int]]) -> None:
    # Write headers
    for col_idx, (key, header, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _WRAP
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Write data rows
    for row_idx, record in enumerate(data, 2):
        for col_idx, (key, _, _) in enumerate(columns, 1):
            value = record.get(key, "")
            if isinstance(value, list):
                value = ", ".join(str(v) for v in value)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = _WRAP

        # Color the status cell
        status = record.get("quality_status", "")
        status_col = next(
            (i for i, (k, _, _) in enumerate(columns, 1) if k == "quality_status"),
            None,
        )
        if status_col and status in _STATUS_FILLS:
            ws.cell(row=row_idx, column=status_col).fill = _STATUS_FILLS[status]

    # Freeze header row
    ws.freeze_panes = "A2"
    # Auto-filter
    ws.auto_filter.ref = ws.dimensions
